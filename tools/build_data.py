# -*- coding: utf-8 -*-
"""
build_data.py — Convertit les classements FISF (XLSX, dossier archives/)
en un fichier de données JavaScript consommé par l'application web.

Usage :
    python tools/build_data.py

Fonctionnement :
    1. Lit tous les fichiers archives/YYYY-MM-DD_*.xlsx, triés par date.
    2. Le plus récent = snapshot courant ; les précédents servent à calculer
       la place précédente (pv) et l'historique des places (bump chart).
    3. Si un seul snapshot est disponible, les places précédentes et
       l'historique sont SIMULÉS (étiquetés simulated=true dans les données).
    4. Extrait l'onglet "Barres et quotas" (quotas par fédération, barres
       par série).
    5. Écrit assets/data/data.js  (window.FISF_DATA = {...})

Pour publier une nouvelle mise à jour du classement :
    déposer le nouveau fichier archives/2026-XX-XX_classements-....xlsx
    puis relancer ce script.
"""

import json
import math
import random
import re
import sys
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
ARCHIVES = ROOT / "archives"
OUT = ROOT / "assets" / "data" / "data.js"

HIST_TOP = 200          # nombre de joueurs dont on conserve l'historique complet
HIST_LEN = 5            # nombre de points d'historique (snapshots)
CODE_VERSION = "2026-09-29b"  # version du code (css/js) : a faire evoluer a chaque modif de code
SERIES_PCT_RE = re.compile(r"^%s([1-6])\s*25-?26$")
BARRE_LABEL_RE = re.compile(r"^(3/4|\d[A-D]?)$")


def norm(s):
    """Normalise un en-tête de colonne pour le rapprochement."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[\s_\.\-]+", "", s).lower()


def find_sheet(wb, keyword):
    for name in wb.sheetnames:
        if keyword in norm(name):
            return wb[name]
    return None


def col_map(headers, wanted):
    """Renvoie {cle: index} en rapprochant les en-têtes normalisés."""
    normed = [norm(h) for h in headers]
    out = {}
    for key, candidates in wanted.items():
        for cand in candidates:
            if cand in normed:
                out[key] = normed.index(cand)
                break
    return out


def num(v, default=None):
    try:
        if v is None or v == "":
            return default
        f = float(v)
        return int(f) if f.is_integer() else round(f, 2)
    except (TypeError, ValueError):
        return default


def cell(row, idx):
    """Valeur de la cellule row[idx], None si index absent."""
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def parse_duplicate(ws):
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    wanted = {
        "lic": ["licence", "license"],
        "nom": ["nom"],
        "pre": ["prenom"],
        "club": ["club"],
        "fed": ["fede", "federation", "fed"],
        "pay": ["pays"],
        "pl": ["place", "placec", "rang"],
        "s25": ["serie2526", "serie25-26"],
        "s26": ["serie2627", "serie26-27", "serie"],
    }
    m = col_map(headers, wanted)
    for i, h in enumerate(headers):
        mm = SERIES_PCT_RE.match(norm(h))
        if mm and norm(h).startswith("%"):
            m.setdefault("s" + mm.group(1), i)
    missing = [k for k in ("lic", "nom", "pl") if k not in m]
    if missing:
        sys.exit(f"[Duplicate] Colonnes introuvables : {missing} — en-têtes : {headers}")
    players = []
    for r in rows:
        if cell(r, m["lic"]) is None:
            continue
        pcts = [num(cell(r, m.get(f"s{i}")), None) for i in range(1, 7)]
        vals = [v for v in pcts if v is not None]
        sc = round(sum(vals) / len(vals), 2) if vals else None
        players.append({
            "lic": num(cell(r, m["lic"])),
            "nom": str(cell(r, m["nom"])).strip(),
            "pre": str(cell(r, m.get("pre"))).strip() if cell(r, m.get("pre")) else "",
            "club": str(cell(r, m.get("club"))).strip() if cell(r, m.get("club")) else "",
            "fed": str(cell(r, m.get("fed"))).strip() if cell(r, m.get("fed")) else "",
            "pay": str(cell(r, m.get("pay"))).strip() if cell(r, m.get("pay")) else "",
            "s25": str(cell(r, m.get("s25"))).strip() if cell(r, m.get("s25")) is not None else "",
            "s26": str(cell(r, m.get("s26"))).strip() if cell(r, m.get("s26")) is not None else "",
            "pl": num(cell(r, m["pl"])),
            "sc": sc,
            "s1": num(cell(r, m.get("s1"))), "s2": num(cell(r, m.get("s2"))),
            "s3": num(cell(r, m.get("s3"))), "s4": num(cell(r, m.get("s4"))),
            "s5": num(cell(r, m.get("s5"))), "s6": num(cell(r, m.get("s6"))),
        })
    players.sort(key=lambda p: (p["pl"] is None, p["pl"]))
    return players


def parse_classic(ws):
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    wanted = {
        "lic": ["licence", "license"],
        "nom": ["nom"],
        "pre": ["prenom"],
        "club": ["club"],
        "fed": ["fede", "federation"],
        "pay": ["pays"],
        "pl": ["place", "rang"],
        "s25": ["serie25-26", "serie2526"],
        "s26": ["serie26-27", "serie2627", "serie"],
        "cote": ["cote"],
        "v": ["victoire", "victoires"],
        "d": ["defaite", "defaites"],
        "n": ["nul", "nuls"],
        "matchs": ["matchs", "matches"],
        "cmin": ["cotemin", "coteminimum"],
        "cmax": ["cotemax", "cotemaximum"],
    }
    m = col_map(headers, wanted)
    missing = [k for k in ("lic", "nom", "pl") if k not in m]
    if missing:
        sys.exit(f"[Classique] Colonnes introuvables : {missing} — en-têtes : {headers}")
    players = []
    for r in rows:
        if cell(r, m["lic"]) is None:
            continue
        players.append({
            "lic": num(cell(r, m["lic"])),
            "nom": str(cell(r, m["nom"])).strip(),
            "pre": str(cell(r, m.get("pre"))).strip() if cell(r, m.get("pre")) else "",
            "club": str(cell(r, m.get("club"))).strip() if cell(r, m.get("club")) else "",
            "fed": str(cell(r, m.get("fed"))).strip() if cell(r, m.get("fed")) else "",
            "pay": str(cell(r, m.get("pay"))).strip() if cell(r, m.get("pay")) else "",
            "s25": str(cell(r, m.get("s25"))).strip() if cell(r, m.get("s25")) is not None else "",
            "s26": str(cell(r, m.get("s26"))).strip() if cell(r, m.get("s26")) is not None else "",
            "pl": num(cell(r, m["pl"])),
            "cote": num(cell(r, m.get("cote"))),
            "v": num(cell(r, m.get("v")), 0), "d": num(cell(r, m.get("d")), 0),
            "n": num(cell(r, m.get("n")), 0), "matchs": num(cell(r, m.get("matchs")), 0),
            "cmin": num(cell(r, m.get("cmin"))), "cmax": num(cell(r, m.get("cmax"))),
        })
    players.sort(key=lambda p: (p["pl"] is None, p["pl"]))
    return players


def parse_barres(ws):
    quotas, barres = {}, []
    seen_barres = set()
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        first = str(cells[0]).strip() if cells and cells[0] is not None else ""
        if first in ("Belgique", "France", "Québec", "Suisse", "Congo Kin",
                     "Sénégal", "Autres", "Totaux"):
            vals = [num(v, 0) for v in cells[1:7]]
            if any(v is not None for v in vals):
                quotas[first] = vals
        for i, c in enumerate(cells[:-1]):
            if isinstance(c, str) and BARRE_LABEL_RE.match(c.strip()):
                v = num(cells[i + 1])
                if v is not None and c.strip() not in seen_barres:
                    seen_barres.add(c.strip())
                    barres.append([c.strip(), v])
    return {"quotas": quotas, "barres": barres}


def load_snapshots():
    files = sorted(ARCHIVES.glob("*.xlsx"))
    if not files:
        sys.exit(f"Aucun fichier XLSX trouvé dans {ARCHIVES}")
    import openpyxl
    snaps = []
    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        dup = find_sheet(wb, "duplicate")
        cla = find_sheet(wb, "classique")
        bar = find_sheet(wb, "barre")
        entry = {"date": f.name[:10], "file": f.name}
        if dup:
            entry["duplicate"] = parse_duplicate(dup)
        if cla:
            entry["classic"] = parse_classic(cla)
        if bar:
            entry["barres"] = parse_barres(bar)
        snaps.append(entry)
        wb.close()
        print(f"  snapshot {f.name} : dup={len(entry.get('duplicate', []))} "
              f"cla={len(entry.get('classic', []))}")
    return snaps


def compute_deltas(players, prev_players):
    """Renseigne pv (place précédente) et nw (nouvel entrant)."""
    prev_by_lic = {p["lic"]: p for p in prev_players}
    for p in players:
        prev = prev_by_lic.get(p["lic"])
        p["pv"] = prev["pl"] if prev else None
        p["nw"] = 0 if prev else 1


def simulate_deltas(players, seed=42):
    """Un seul snapshot : simule des places précédentes plausibles (démo)."""
    rng = random.Random(seed)
    for p in players:
        if rng.random() < 0.015:
            p["pv"], p["nw"] = None, 1
        else:
            spread = 4 if p["pl"] <= 100 else (15 if p["pl"] <= 1000 else 60)
            pv = p["pl"] + rng.randint(-spread, spread)
            p["pv"], p["nw"] = max(1, pv), 0


def build_history(players, snaps_players, dates):
    """Historique des places pour le top HIST_TOP.
    snaps_players : liste des listes de joueurs (ordre chronologique)."""
    hist = {}
    top = players[:HIST_TOP]
    if len(snaps_players) >= 2:
        for p in top:
            series = []
            for snap in snaps_players:
                match = next((q for q in snap if q["lic"] == p["lic"]), None)
                series.append(match["pl"] if match else None)
            hist[str(p["lic"])] = series
    else:
        rng = random.Random(7)
        for p in top:
            pts = [p["pl"]]
            cur = p["pl"]
            for _ in range(HIST_LEN - 1):
                cur = max(1, cur + rng.randint(-4, 8))
                pts.append(cur)
            hist[str(p["lic"])] = list(reversed(pts))
    return hist


def main():
    print("Lecture des snapshots…")
    snaps = load_snapshots()
    current = snaps[-1]
    only_one = len(snaps) == 1
    dates = [s["date"] for s in snaps]
    if only_one:
        y, m, d = map(int, dates[0].split("-"))
        while len(dates) < HIST_LEN:
            m -= 1
            if m == 0:
                m, y = 12, y - 1
            d = min(d, 28)
            dates.insert(0, f"{y:04d}-{m:02d}-{d:02d}")
    data = {
        "season": "2026-2027",
        "generated": current["date"],
        "snapshots": dates[-HIST_LEN:],
        "simulated": only_one,
        "barres": current.get("barres", {"quotas": {}, "barres": []}),
        "disciplines": {},
    }
    for disc, key in (("duplicate", "duplicate"), ("classic", "classic")):
        players = current.get(key)
        if not players:
            continue
        if only_one:
            simulate_deltas(players, seed=42 if disc == "duplicate" else 43)
            hist = build_history(players, [], dates)
        else:
            prev_snaps = [s[key] for s in snaps[:-1] if key in s]
            if prev_snaps:
                compute_deltas(players, prev_snaps[-1])
            else:
                simulate_deltas(players, seed=42 if disc == "duplicate" else 43)
            hist_src = [s[key] for s in snaps if key in s]
            hist = build_history(players, hist_src, dates)
        countries = {}
        series = {}
        for p in players:
            countries[p["pay"]] = countries.get(p["pay"], 0) + 1
            series[p["s26"]] = series.get(p["s26"], 0) + 1
        data["disciplines"][disc] = {
            "players": players,
            "hist": hist,
            "countries": dict(sorted(countries.items(), key=lambda kv: -kv[1])),
            "series": dict(sorted(series.items(), key=lambda kv: -kv[1])),
        }
        print(f"{disc} : {len(players)} joueurs, "
              f"{len(data['disciplines'][disc]['hist'])} historiques")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    OUT.write_text("window.FISF_DATA=" + payload + ";\n", encoding="utf-8")
    mb = OUT.stat().st_size / 1e6
    print(f"Écrit : {OUT} ({mb:.2f} Mo)")
    # Cache-busting : horodate les références d'assets dans index.html pour que
    # les navigateurs rechargent css/js/data après chaque modification.
    # - data.js : version = date du snapshot (regénéré à chaque build)
    # - style.css / app.js : version = CODE_VERSION (à faire évoluer à chaque
    #   modification de code sans régénération de données)
    idx = ROOT / "index.html"
    html = idx.read_text(encoding="utf-8")
    html = re.sub(r"(assets/data/data\.js\?v=)[0-9]{4}-[0-9]{2}-[0-9]{2}[a-z]?",
                  r"\g<1>" + current["date"], html)
    html = re.sub(r"((?:assets/css/style\.css|assets/js/app\.js)\?v=)[0-9]{4}-[0-9]{2}-[0-9]{2}[a-z]?",
                  r"\g<1>" + CODE_VERSION, html)
    idx.write_text(html, encoding="utf-8")
    print(f"index.html mis a jour (data v={current['date']}, code v={CODE_VERSION})")
    if only_one:
        print("NOTE : un seul snapshot -> places precedentes et historique SIMULES "
              "(data.simulated=true). Déposez les classements passés dans archives/ "
              "puis relancez pour obtenir les vraies évolutions.")


if __name__ == "__main__":
    main()
